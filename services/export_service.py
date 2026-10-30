"""
services/export_service.py
ส่งออกรายงานข้อมูลร้านเป็นไฟล์ Microsoft Excel (.xlsx) และสำรองรายเดือนอัตโนมัติ (OWN-5, D5, D8, E7, ข้อ 13.2):
- เจ้าของร้านเลือกช่วงเวลาส่งออกได้สูงสุด 1 ปี (รวมถึงปีปฏิทิน 1 ม.ค. – 31 ธ.ค., 12 เดือนล่าสุด, รายเดือน หรือช่วงกำหนดเอง)
- ประกอบด้วย 6 ชีตชัดเจน: สรุปภาพรวม, ออเดอร์ทั้งหมด, รายการสินค้าที่ซื้อ, สรุปลูกค้า, รายจ่ายร้าน, สต็อกคงเหลือ
- จัดหน้า Excel ให้สวยงาม มีสีแบ่ง รายรับ (โทนเขียวอ่อน) / รายจ่าย (โทนแดงอ่อน) ชัดเจน พร้อมหัวตารางเข้ม ความกว้างคอลัมน์ รูปแบบสกุลเงิน และแถวสรุปรวม
- แยกไฟล์ .xlsx รายงานออกจากไฟล์สำรองฐานข้อมูล SQLite (.db) อย่างชัดเจน และไม่เขียนทับข้อมูลฐานข้อมูล
"""
import calendar
from datetime import date, datetime, timedelta
from pathlib import Path

import config
from core import permissions
from core.constants import ORDER_STATUS_LABELS
from database import repo_expenses, repo_orders, repo_products, repo_settings
from services import audit_service, settings_service

MAX_EXPORT_DAYS = 366


def validate_date_range(start_date: str, end_date: str) -> tuple[str, str]:
    """
    ตรวจสอบช่วงวันที่ (รูปแบบ YYYY-MM-DD):
    - start_date <= end_date
    - ช่วงห่างไม่เกิน 1 ปี (366 วัน) ตามข้อกำหนด PDF
    """
    try:
        s_dt = datetime.strptime(str(start_date or "").strip()[:10], "%Y-%m-%d").date()
        e_dt = datetime.strptime(str(end_date or "").strip()[:10], "%Y-%m-%d").date()
    except ValueError:
        raise ValueError("รูปแบบวันที่ไม่ถูกต้อง (กรุณาใช้ ปปปป-ดด-วว เช่น 2026-01-01)") from None

    if s_dt > e_dt:
        raise ValueError("วันเริ่มต้นต้องไม่มากกว่าวันสิ้นสุด")
    days_span = (e_dt - s_dt).days + 1
    if days_span > MAX_EXPORT_DAYS:
        raise ValueError(f"เลือกช่วงวันที่ได้สูงสุดไม่เกิน 1 ปี ({MAX_EXPORT_DAYS} วัน)")
    return s_dt.isoformat(), e_dt.isoformat()


def range_calendar_year(year: int = None) -> tuple[str, str]:
    """ช่วงวันที่ทั้งปีปฏิทิน 1 ม.ค. – 31 ธ.ค. ของปีที่เลือก"""
    y = int(year or date.today().year)
    return f"{y:04d}-01-01", f"{y:04d}-12-31"


def range_last_12_months(ref_date: date = None) -> tuple[str, str]:
    """ช่วง 1 ปี (365 วันล่าสุด จนถึงวันนี้)"""
    end_d = ref_date or date.today()
    start_d = end_d - timedelta(days=364)
    return start_d.isoformat(), end_d.isoformat()


def range_month(year: int, month: int) -> tuple[str, str]:
    """ช่วงวันที่ของเดือนนั้น (วันแรก ถึง วันสุดท้ายของเดือน)"""
    y, m = int(year), int(month)
    if not 1 <= m <= 12:
        raise ValueError("เดือนต้องอยู่ระหว่าง 1 ถึง 12")
    last_day = calendar.monthrange(y, m)[1]
    return f"{y:04d}-{m:02d}-01", f"{y:04d}-{m:02d}-{last_day:02d}"


def build_export_data(start_date: str, end_date: str, include_personal: bool = True) -> dict:
    """
    รวบรวมข้อมูลจริงจากฐานข้อมูลสำหรับช่วงวันที่ที่กำหนด (ใช้ร่วมกันทั้ง .xlsx และ Google Sheets)
    """
    s_date, e_date = validate_date_range(start_date, end_date)

    revenue = repo_orders.sum_revenue(s_date, e_date)
    expense_total = repo_expenses.sum_expenses(s_date, e_date)
    orders = [dict(r) for r in repo_orders.list_orders_in_range(s_date, e_date)]
    sold_items = [dict(r) for r in repo_orders.list_sold_items_in_range(s_date, e_date)]
    customers = [dict(r) for r in repo_orders.list_customer_summaries_in_range(s_date, e_date)]
    expenses = [dict(r) for r in repo_expenses.list_expenses(s_date, e_date)]
    stock_rows = [dict(r) for r in repo_products.list_products({"active_only": False, "include_deleted": False})]

    cogs_total = round(sum(float(i.get("unit_cost") or 0) * int(i.get("qty") or 0) for i in sold_items), 2)
    paid_orders_count = sum(1 for o in orders if o["status"] in repo_orders.REVENUE_STATUSES)
    cancelled_orders_count = sum(1 for o in orders if o["status"] == "cancelled")
    net_profit = round(revenue - expense_total - cogs_total, 2)
    pre_vat_rev = round(revenue / 1.07, 2)
    vat_rev = round(revenue - pre_vat_rev, 2)

    summary_rows = [
        ["หัวข้อสรุป", "ค่า / จำนวน", "หมายเหตุ"],
        ["ช่วงวันที่รายงาน", f"{s_date} ถึง {e_date}", "นับรวมทั้งวันเริ่มต้นและวันสิ้นสุด (00:00:00 - 23:59:59)"],
        ["[รายรับ] รายรับรวมสุทธิ (บาท)", revenue, "นับเฉพาะออเดอร์ที่ชำระเงินแล้ว (paid, shipped, completed)"],
        ["[รายรับ] มูลค่าขายก่อนภาษี (Pre-VAT)", pre_vat_rev, "คำนวณถอดภาษีมูลค่าเพิ่ม 7%"],
        ["[รายรับ] ภาษีมูลค่าเพิ่ม (VAT 7%)", vat_rev, "ภาษีขาย 7% ที่รวมอยู่ในยอดชำระแล้ว"],
        ["[รายจ่าย] รายจ่ายร้านที่บันทึก (บาท)", expense_total, "จากตารางรายจ่าย (expenses)"],
        ["[รายจ่าย] ต้นทุนสินค้าที่ขาย (บาท)", cogs_total, "คำนวณจาก cost_price ของสินค้า (ถ้าระบุ)"],
        ["[สรุป] กำไรสุทธิเบื้องต้น (บาท)", net_profit, "รายรับรวม - รายจ่ายร้าน - ต้นทุนสินค้า"],
        ["จำนวนคำสั่งซื้อทั้งหมด (รายการ)", len(orders), "รวมทุกสถานะในช่วงเวลา"],
        ["จำนวนคำสั่งซื้อที่ชำระเงินแล้ว (รายการ)", paid_orders_count, "สถานะ paid, shipped, completed"],
        ["จำนวนคำสั่งซื้อที่ยกเลิก (รายการ)", cancelled_orders_count, "สถานะ cancelled"],
        ["จำนวนลูกค้าที่ซื้อสำเร็จ (คน)", len(customers), "นับจำนวนบัญชีลูกค้าไม่ซ้ำ"],
        ["จำนวนชิ้นสินค้าที่ขายได้รวม (ชิ้น)", sum(int(i["qty"]) for i in sold_items), "เฉพาะออเดอร์ที่ชำระเงินแล้ว"],
    ]

    order_rows = [
        ["เลขที่ออเดอร์", "วันที่สั่งซื้อ", "ชื่อลูกค้า", "เบอร์ผู้รับ", "ที่อยู่จัดส่ง",
         "ยอดสินค้า (บาท)", "ส่วนลด (บาท)", "ค่าส่ง (บาท)", "ยอดสุทธิ (บาท)", "สถานะ"]
    ]
    for o in orders:
        order_rows.append([
            o["id"],
            o["created_at"],
            o["buyer_name"],
            o["recipient_phone"] if include_personal else "***",
            o["shipping_address"] if include_personal else "***",
            float(o["subtotal"]),
            float(o["discount"]),
            float(o["shipping_fee"]),
            float(o["total"]),
            ORDER_STATUS_LABELS.get(o["status"], o["status"]),
        ])
    if orders:
        order_rows.append([
            "รวมทั้งหมด", f"{len(orders)} ออเดอร์", "", "", "",
            round(sum(float(o["subtotal"]) for o in orders), 2),
            round(sum(float(o["discount"]) for o in orders), 2),
            round(sum(float(o["shipping_fee"]) for o in orders), 2),
            round(sum(float(o["total"]) for o in orders), 2),
            f"ชำระแล้ว {paid_orders_count} รายการ",
        ])

    item_rows = [
        ["วันที่สั่งซื้อ", "เลขที่ออเดอร์", "ชื่อลูกค้า", "รหัสสินค้า", "ชื่อสินค้า",
         "ตัวเลือกสินค้า", "จำนวน (ชิ้น)", "ราคาต่อชิ้น (บาท)", "ยอดรวมรายการ (บาท)"]
    ]
    for it in sold_items:
        item_rows.append([
            it["order_date"],
            it["order_id"],
            it["buyer_name"],
            it["product_id"],
            it["product_name"],
            it["option_label"] or "-",
            int(it["qty"]),
            float(it["price_at_purchase"]),
            float(it["line_total"]),
        ])
    if sold_items:
        item_rows.append([
            "รวมยอดขายสินค้า", "", "", "", "", "",
            sum(int(it["qty"]) for it in sold_items),
            "",
            round(sum(float(it["line_total"]) for it in sold_items), 2),
        ])

    customer_rows = [
        ["รหัสลูกค้า", "ชื่อ-นามสกุล", "อีเมล", "เบอร์โทรศัพท์", "จำนวนออเดอร์ที่ซื้อ", "ยอดซื้อสะสมรวม (บาท)"]
    ]
    for c in customers:
        customer_rows.append([
            c["user_id"],
            c["full_name"],
            c["email"] if include_personal else "***",
            (c["phone"] or "-") if include_personal else "***",
            int(c["order_count"]),
            float(c["total_spent"]),
        ])
    if customers:
        customer_rows.append([
            "รวมลูกค้าทั้งหมด", f"{len(customers)} คน", "", "",
            sum(int(c["order_count"]) for c in customers),
            round(sum(float(c["total_spent"]) for c in customers), 2),
        ])

    expense_rows = [
        ["รหัสรายจ่าย", "วันที่จ่าย", "หมวดหมู่รายจ่าย", "จำนวนเงิน (บาท)", "หมายเหตุ", "ผู้บันทึก"]
    ]
    for ex in expenses:
        expense_rows.append([
            ex["id"],
            ex["spent_on"],
            ex["category"],
            float(ex["amount"]),
            ex["note"] or "-",
            ex.get("created_by_name") or "-",
        ])
    if expenses:
        expense_rows.append([
            "รวมรายจ่ายทั้งหมด", "", f"{len(expenses)} รายการ",
            round(sum(float(ex["amount"]) for ex in expenses), 2),
            "", "",
        ])

    stock_sheet_rows = [
        ["รหัสสินค้า", "ชื่อสินค้า", "หมวดหมู่", "แบรนด์", "ราคาปกติ (บาท)", "ราคาขายจริง (บาท)",
         "สต็อกคงเหลือ (ชิ้น)", "สถานะการขาย"]
    ]
    for p in stock_rows:
        stock_sheet_rows.append([
            p["id"],
            p["name"],
            p.get("category_name") or "-",
            p.get("brand_name") or "-",
            float(p["price"]),
            float(p["effective_price"]),
            int(p["stock"] or 0),
            "เปิดขาย" if p["is_active"] else "ปิดขาย",
        ])

    return {
        "start_date": s_date,
        "end_date": e_date,
        "revenue": revenue,
        "expense_total": expense_total,
        "net_profit": net_profit,
        "orders_count": len(orders),
        "customers_count": len(customers),
        "summary": {
            "gross_revenue": revenue,
            "expense_total": expense_total,
            "net_profit": net_profit,
            "unique_customers": len(customers),
            "paid_orders_count": paid_orders_count,
        },
        "sheets": {
            "สรุปภาพรวม": summary_rows,
            "ออเดอร์ทั้งหมด": order_rows,
            "รายการสินค้าที่ซื้อ": item_rows,
            "สรุปลูกค้า": customer_rows,
            "รายจ่ายร้าน": expense_rows,
            "สต็อกคงเหลือ": stock_sheet_rows,
        },
    }


def _write_xlsx(data: dict, target_path: Path) -> Path:
    """
    เขียนไฟล์ .xlsx พร้อมจัดหน้าและแยกสี รายรับ (เขียวอ่อน) / รายจ่าย (แดงอ่อน) ชัดเจน (ข้อ 13.2)
    """
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # โทนสีหัวตารางและสีแยกรายรับ/รายจ่าย
    header_dark_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_rev_fill = PatternFill(start_color="14532D", end_color="14532D", fill_type="solid")
    header_exp_fill = PatternFill(start_color="7F1D1D", end_color="7F1D1D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    rev_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")       # เขียวอ่อน (รายรับ)
    rev_alt_fill = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")
    rev_font = Font(name="Calibri", size=11, bold=True, color="166534")

    exp_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")       # แดงอ่อน (รายจ่าย)
    exp_alt_fill = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")
    exp_font = Font(name="Calibri", size=11, bold=True, color="991B1B")

    profit_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")    # ทองอ่อน (กำไรสุทธิ/แถวรวม)
    profit_font = Font(name="Calibri", size=11, bold=True, color="92400E")

    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    body_font = Font(name="Calibri", size=11, color="1E293B")
    bold_body_font = Font(name="Calibri", size=11, bold=True, color="0F172A")

    thin_side = Side(border_style="thin", color="CBD5E1")
    double_bottom = Side(border_style="double", color="1E293B")
    top_thick = Side(border_style="medium", color="1E293B")
    cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    total_border = Border(left=thin_side, right=thin_side, top=top_thick, bottom=double_bottom)

    for sheet_title, rows in data["sheets"].items():
        ws = wb.create_sheet(title=sheet_title)
        is_expense_sheet = (sheet_title == "รายจ่ายร้าน")
        is_revenue_sheet = sheet_title in ("ออเดอร์ทั้งหมด", "รายการสินค้าที่ซื้อ", "สรุปลูกค้า")
        is_summary_sheet = (sheet_title == "สรุปภาพรวม")

        total_rows = len(rows)
        for r_idx, row_values in enumerate(rows, start=1):
            ws.append(row_values)
            ws.row_dimensions[r_idx].height = 26 if r_idx == 1 else 22
            first_text = str(row_values[0] or "")
            is_total_row = (r_idx > 1 and r_idx == total_rows and first_text.startswith("รวม"))

            for c_idx, val in enumerate(row_values, start=1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.border = total_border if is_total_row else cell_border

                if r_idx == 1:
                    cell.fill = header_exp_fill if is_expense_sheet else (header_rev_fill if is_revenue_sheet else header_dark_fill)
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    continue

                # จัดรูปแบบตัวเลข/สกุลเงิน
                if isinstance(val, float):
                    cell.number_format = "#,##0.00"
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif isinstance(val, int):
                    cell.number_format = "#,##0"
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

                # แยกสี รายรับ / รายจ่าย / กำไรสุทธิ
                if is_summary_sheet:
                    if "[รายรับ]" in first_text or "รายรับ" in first_text:
                        cell.fill = rev_fill
                        cell.font = rev_font if c_idx in (1, 2) else body_font
                    elif "[รายจ่าย]" in first_text or "รายจ่าย" in first_text or "ต้นทุน" in first_text:
                        cell.fill = exp_fill
                        cell.font = exp_font if c_idx in (1, 2) else body_font
                    elif "กำไรสุทธิ" in first_text:
                        cell.fill = profit_fill
                        cell.font = profit_font if float(data.get("net_profit", 0)) >= 0 else exp_font
                    else:
                        if r_idx % 2 == 1:
                            cell.fill = zebra_fill
                        cell.font = bold_body_font if c_idx == 1 else body_font
                elif is_total_row:
                    cell.fill = exp_fill if is_expense_sheet else (rev_fill if is_revenue_sheet else profit_fill)
                    cell.font = exp_font if is_expense_sheet else (rev_font if is_revenue_sheet else profit_font)
                elif is_expense_sheet:
                    cell.fill = exp_fill if c_idx == 4 else (exp_alt_fill if r_idx % 2 == 0 else PatternFill(fill_type=None))
                    cell.font = exp_font if c_idx == 4 else body_font
                elif is_revenue_sheet:
                    # ไฮไลต์คอลัมน์ยอดเงินรายรับด้วยสีเขียวอ่อน
                    is_money_col = isinstance(val, float) and c_idx == len(row_values) or (sheet_title == "ออเดอร์ทั้งหมด" and c_idx == 9)
                    if is_money_col:
                        cell.fill = rev_fill
                        cell.font = rev_font
                    else:
                        if r_idx % 2 == 0:
                            cell.fill = rev_alt_fill
                        cell.font = body_font
                else:
                    if r_idx % 2 == 0:
                        cell.fill = zebra_fill
                    cell.font = body_font

        ws.freeze_panes = "A2"
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = max((len(str(c.value or "")) for c in col), default=12)
            ws.column_dimensions[col_letter].width = min(max(max_len + 6, 15), 50)

    target_path = Path(target_path)
    target_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target_path)
    return target_path


def export_range(start_date: str, end_date: str, dest_path=None, require_auth: bool = True) -> Path:
    """
    ส่งออกรายงาน .xlsx สำหรับช่วงวันที่ที่กำหนด (เจ้าของร้านเท่านั้น)
    ห้ามเขียนทับไฟล์ฐานข้อมูล .db เด็ดขาด
    """
    if require_auth:
        permissions.require_any(permissions.BACKUP_MANAGE, permissions.DASHBOARD_REVENUE)

    data = build_export_data(start_date, end_date, include_personal=True)
    s_date, e_date = data["start_date"], data["end_date"]

    if dest_path is None:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        out_path = Path(config.EXPORTS_DIR) / f"report_{s_date}_to_{e_date}_{stamp}.xlsx"
    else:
        out_path = Path(dest_path)
        if out_path.suffix.lower() != ".xlsx":
            out_path = out_path.with_suffix(".xlsx")

    if out_path.resolve() == Path(config.DB_PATH).resolve():
        raise ValueError("ไม่สามารถบันทึกรายงานทับไฟล์ฐานข้อมูลได้")

    saved = _write_xlsx(data, out_path)
    if require_auth:
        audit_service.log("export.xlsx", "report", None, {"start_date": s_date, "end_date": e_date, "file": saved.name})
    return saved


def set_auto_export_folder(folder_path: str) -> str:
    """เจ้าของร้านเลือก/ยืนยันโฟลเดอร์ปลายทางสำหรับบันทึกรายงานรายเดือนอัตโนมัติ"""
    permissions.require(permissions.BACKUP_MANAGE)
    folder = str(folder_path or "").strip()
    if not folder:
        repo_settings.set("auto_export_dir", "")
        return ""
    p = Path(folder).resolve()
    p.mkdir(parents=True, exist_ok=True)
    repo_settings.set("auto_export_dir", str(p))
    audit_service.log("export.set_auto_folder", "settings", None, {"folder": str(p)})
    return str(p)


def auto_export_previous_month(force: bool = False) -> dict:
    """
    ส่งออกรายงานประจำเดือนก่อนหน้าโดยอัตโนมัติลงโฟลเดอร์ที่เจ้าของร้านยืนยันไว้ (หรือ data/exports/monthly เมื่อสั่งรัน)
    """
    today = date.today()
    first_this_month = today.replace(day=1)
    prev_month_last_day = first_this_month - timedelta(days=1)
    ym = prev_month_last_day.strftime("%Y-%m")

    last_done = (settings_service.get("last_auto_export") or "").strip()
    if not force and last_done >= ym:
        return {"exported": False, "month": ym, "path": None, "reason": "ส่งออกเดือนนี้ไปแล้ว"}

    configured_dir = (settings_service.get("auto_export_dir") or "").strip()
    if not configured_dir and not force:
        return {"exported": False, "month": ym, "path": None, "reason": "ยังไม่ได้ตั้งค่าโฟลเดอร์ส่งออกอัตโนมัติ"}

    target_dir = Path(configured_dir) if configured_dir else Path(config.MONTHLY_EXPORTS_DIR)
    target_dir.mkdir(parents=True, exist_ok=True)
    s_date, e_date = range_month(prev_month_last_day.year, prev_month_last_day.month)
    out_file = target_dir / f"BadmintonProShop_Report_{ym}.xlsx"

    saved = export_range(s_date, e_date, dest_path=out_file, require_auth=False)
    repo_settings.set("last_auto_export", ym)
    return {"exported": True, "month": ym, "path": str(saved), "reason": "สำเร็จ"}


def list_exported_files() -> list[dict]:
    """แสดงรายการไฟล์ .xlsx ที่เคยส่งออกไว้ในโฟลเดอร์ของระบบและโฟลเดอร์อัตโนมัติ"""
    permissions.require_any(permissions.BACKUP_MANAGE, permissions.DASHBOARD_REVENUE)
    folders = {Path(config.EXPORTS_DIR), Path(config.MONTHLY_EXPORTS_DIR)}
    custom_dir = (settings_service.get("auto_export_dir") or "").strip()
    if custom_dir and Path(custom_dir).is_dir():
        folders.add(Path(custom_dir))

    files = []
    seen = set()
    for folder in folders:
        if not folder.is_dir():
            continue
        for f in folder.glob("*.xlsx"):
            resolved = f.resolve()
            if resolved in seen:
                continue
            seen.add(resolved)
            try:
                st = f.stat()
                files.append({
                    "name": f.name,
                    "path": str(resolved),
                    "size": st.st_size,
                    "created_at": datetime.fromtimestamp(st.st_mtime).strftime("%Y-%m-%d %H:%M:%S"),
                })
            except OSError:
                continue
    files.sort(key=lambda x: x["created_at"], reverse=True)
    return files


def sync_to_google_sheets(start_date: str, end_date: str, include_personal: bool = False) -> dict:
    """
    ส่งออกข้อมูลไปยัง Google Sheets (ทางเลือกเสริม E7 — import gspread แบบ lazy)
    """
    permissions.require(permissions.BACKUP_MANAGE)
    data = build_export_data(start_date, end_date, include_personal=include_personal)

    if not config.GOOGLE_SERVICE_ACCOUNT_FILE or not config.GOOGLE_SHEET_ID:
        raise ValueError("ยังไม่ได้ตั้งค่า GOOGLE_SERVICE_ACCOUNT_FILE และ GOOGLE_SHEET_ID ในไฟล์ .env")

    try:
        import gspread
    except ImportError:
        raise ValueError("ยังไม่ได้ติดตั้งส่วนเสริม Google Sheets (gspread)") from None

    sa_path = Path(config.GOOGLE_SERVICE_ACCOUNT_FILE)
    if not sa_path.is_file():
        raise ValueError("ไม่พบไฟล์กุญแจ Service Account สำหรับ Google Sheets")

    client = gspread.service_account(filename=str(sa_path))
    sh = client.open_by_key(config.GOOGLE_SHEET_ID)
    tab_name = f"Report_{data['start_date']}_{data['end_date']}"[:95]
    try:
        ws = sh.worksheet(tab_name)
        ws.clear()
    except Exception:
        ws = sh.add_worksheet(title=tab_name, rows="500", cols="20")

    combined_rows = []
    for title, sheet_rows in data["sheets"].items():
        combined_rows.append([f"=== {title} ==="])
        combined_rows.extend([[str(c) for c in r] for r in sheet_rows])
        combined_rows.append([])
    ws.update("A1", combined_rows)
    audit_service.log(
        "export.google_sheets", "report", None,
        {"start_date": data["start_date"], "end_date": data["end_date"], "include_personal": bool(include_personal)},
    )
    return {"tab_name": tab_name, "rows_written": len(combined_rows), "include_personal": bool(include_personal)}


def export_report_xlsx(start_date: str, end_date: str, dest_path=None) -> dict:
    saved_path = export_range(start_date, end_date, dest_path=dest_path, require_auth=True)
    data = build_export_data(start_date, end_date, include_personal=False)
    summary = data.get("summary") or {}
    return {
        "path": str(saved_path),
        "gross_revenue": summary.get("gross_revenue", 0.0),
        "net_profit": summary.get("net_profit", 0.0),
        "customer_count": summary.get("unique_customers", 0),
        "order_count": summary.get("paid_orders_count", 0),
    }


def configure_auto_monthly_export(enabled: bool, folder_path: str) -> dict:
    saved_folder = set_auto_export_folder(folder_path if enabled else "")
    repo_settings.set("export_auto_monthly_enabled", "1" if enabled else "0")
    repo_settings.set("export_auto_folder", saved_folder)
    return {"enabled": bool(enabled), "folder": saved_folder}
